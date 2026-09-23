"""Independently cross-check submitted tables against the official workbooks.

AI-assisted implementation: OpenAI Codex (GPT-6, OpenAI), used 2026-09-23.
The model release date was not available in this environment and needs team review.
"""

from __future__ import annotations

from collections import Counter
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
INPUT = HERE.parent / "huawei"
RUN = HERE / "runs" / "reproduced"
TABLES = RUN / "tmp" / "tables"
CHECKS = []


def check(name: str, passed: bool, detail: object = "") -> None:
    CHECKS.append({"check": name, "passed": bool(passed), "detail": str(detail)})


def sheet_rows(path: Path, sheet: str):
    book = load_workbook(path, read_only=True, data_only=True)
    rows = [[cell.value for cell in row] for row in book[sheet].iter_rows()]
    book.close()
    return rows


def csv(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / f"{name}.csv")


def box_ids(text: str) -> list[str]:
    return [value.strip() for value in str(text).split(",") if value.strip()]


def no_overlap(frame: pd.DataFrame) -> bool:
    for _, group in frame.groupby("资源编号"):
        group = group.sort_values(["开始时刻（s）", "结束时刻（s）"])
        if any(group["开始时刻（s）"].iloc[i] < group["结束时刻（s）"].iloc[i - 1] - 1e-6
               for i in range(1, len(group))):
            return False
    return True


def main() -> int:
    base = INPUT / "data" / "无人机应急物资运输基础数据"
    raw_boxes = sheet_rows(base / "物资需求与配送时限.xlsx", "逐箱货箱清单")[1:]
    official = {str(row[0]): row for row in raw_boxes if row[0]}
    transport_rows = sheet_rows(base / "运输无人机数据.xlsx", "数据")
    models = {str(row[0]): row for row in transport_rows[2:5]}
    relay_rows = sheet_rows(base / "中继无人机数据.xlsx", "数据")
    inventory = {f"UAV_{model}": sum(row[1] == model for row in transport_rows[8:16])
                 for model in "ABC"}
    inventory.update({f"BAT_{row[0]}": int(row[1]) for row in transport_rows[19:22]})
    inventory["RELAY"] = sum(row[0] == "R01" or row[0] == "R02" for row in relay_rows[6:8])
    inventory["RELAY_ENERGY"] = int(relay_rows[11][1])
    check("official_box_count", len(official) == 80, len(official))

    q1 = csv("q1_batching")
    q1_assignments = []
    q1_ok = True
    for _, trip in q1.iterrows():
        ids = box_ids(trip["货箱编号列表"])
        q1_assignments.extend(ids)
        if any(box not in official for box in ids):
            q1_ok = False
            continue
        model = models[str(trip["机型编号"])]
        q1_ok &= all(official[box][1] == trip["服务区编号"] for box in ids)
        q1_ok &= abs(sum(float(official[box][3]) for box in ids) - trip["总质量（kg）"]) < 1e-6
        q1_ok &= abs(sum(float(official[box][4]) for box in ids) - trip["总体积（m³）"]) < 1e-8
        q1_ok &= trip["总质量（kg）"] <= float(model[3]) + 1e-6
        q1_ok &= trip["总体积（m³）"] <= float(model[4]) + 1e-8
        q1_ok &= trip["架次能耗（kWh）"] <= float(model[8]) * (1 - float(model[9]) / 100) + 1e-6
        q1_ok &= abs(trip["返航SOC（%）"] - 100 * (1 - trip["架次能耗（kWh）"] / float(model[8]))) < 1e-5
    check("q1_every_box_once", Counter(q1_assignments) == Counter(official.keys()), len(q1_assignments))
    check("q1_official_mass_volume_energy", q1_ok, len(q1))

    for question in ("q2", "q3"):
        deliveries = csv(f"{question}_box_delivery")
        trips = csv("q2_transport_sorties" if question == "q2" else "q3_transport_schedule")
        resources = csv(f"{question}_resource_schedule")
        trip_map = {str(row["架次编号"]): row for _, row in trips.iterrows()}
        ids = list(deliveries["货箱编号"].astype(str))
        check(f"{question}_every_box_once", Counter(ids) == Counter(official.keys()), len(ids))
        delivered_by_trip = {}
        deadline_ok = True
        for _, row in deliveries.iterrows():
            box = str(row["货箱编号"])
            trip_id = str(row["架次编号"])
            if box not in official or trip_id not in trip_map:
                deadline_ok = False
                continue
            delivered_by_trip.setdefault(trip_id, []).append(box)
            source = official[box]
            hard = float(source[7]) if source[2] == "医疗物资" else (
                float(source[6]) if source[5] == "是" else None)
            deadline_ok &= source[1] == row["服务区编号"]
            deadline_ok &= hard is None or float(row["交付完成时刻（s）"]) <= hard + 1e-6
            deadline_ok &= float(row["交付完成时刻（s）"]) <= float(trip_map[trip_id]["返回O01时刻（s）"]) + 1e-6
        check(f"{question}_official_hard_deadlines", deadline_ok)
        trip_ok = True
        for trip_id, trip in trip_map.items():
            assigned = box_ids(trip["货箱编号列表"])
            if any(box not in official for box in assigned):
                trip_ok = False
                continue
            model = models[str(trip["机型编号"])]
            trip_ok &= Counter(assigned) == Counter(delivered_by_trip.get(trip_id, []))
            trip_ok &= sum(float(official[box][3]) for box in assigned) <= float(model[3]) + 1e-6
            trip_ok &= sum(float(official[box][4]) for box in assigned) <= float(model[4]) + 1e-8
            trip_ok &= float(trip["架次能耗（kWh）"]) <= float(model[8]) * (1 - float(model[9]) / 100) + 1e-6
            trip_ok &= abs(float(trip["返航SOC（%）"]) - 100 * (1 - float(trip["架次能耗（kWh）"]) / float(model[8]))) < 1e-5
            trip_ok &= trip["开始时刻（s）"] <= trip["返回O01时刻（s）"]
        check(f"{question}_trip_capacity_energy_consistency", trip_ok, len(trips))
        check(f"{question}_resource_intervals", no_overlap(resources), len(resources))

    relays = csv("q3_relay_sorties")
    relay_map = {str(row["中继架次编号"]): row for _, row in relays.iterrows()}
    relay_ok = all(
        row["架次能耗（kWh）"] <= float(relay_rows[2][7]) * (1 - float(relay_rows[2][8]) / 100) + 1e-6
        and row["返航SOC（%）"] >= float(relay_rows[2][8]) - 1e-6
        for _, row in relays.iterrows())
    check("q3_relay_energy", relay_ok, len(relays))
    intervals = csv("q3_communication_coverage")
    coverage_ok = set(intervals["运输架次编号"]) == set(csv("q3_transport_schedule")["架次编号"])
    max_gap = 0.0
    for _, group in intervals.groupby("运输架次编号"):
        group = group.sort_values("开始时刻（s）")
        for i, (_, row) in enumerate(group.iterrows()):
            coverage_ok &= row["保障方式"] in ("DIRECT", "RELAY")
            coverage_ok &= row["结束时刻（s）"] >= row["开始时刻（s）"] - 1e-6
            if i:
                gap = row["开始时刻（s）"] - group.iloc[i - 1]["结束时刻（s）"]
                max_gap = max(max_gap, gap)
                coverage_ok &= -1e-6 <= gap <= 1.01
            if row["保障方式"] == "RELAY":
                relay_id = str(row["中继架次编号"])
                coverage_ok &= relay_id in relay_map
                if relay_id in relay_map:
                    relay = relay_map[relay_id]
                    coverage_ok &= row["开始时刻（s）"] >= relay["建链完成时刻（s）"] - 1e-6
                    coverage_ok &= row["结束时刻（s）"] <= relay["服务结束时刻（s）"] + 1e-6
    check("q3_exported_communication_intervals", coverage_ok, f"max_gap_s={max_gap:.6f}")

    q4_summary = csv("q4_summary")
    for k in (2, 3):
        groups = csv(f"q4_partition_K{k}")
        assignment = {}
        for _, row in groups.iterrows():
            for service in box_ids(row["服务区列表"]):
                assignment.setdefault(service, []).append(row["任务组编号"])
        services = {str(box[1]) for box in official.values()}
        routes_ok = all(len({assignment.get(stop, [None])[0] for stop in
                             str(trip["访问服务区顺序"]).split("->")}) == 1
                        for _, trip in csv("q3_transport_schedule").iterrows())
        check(f"q4_k{k}_partition_and_route", set(assignment) == services and
              all(len(labels) == 1 for labels in assignment.values()) and
              len(groups) == k and routes_ok)
        summary = q4_summary[q4_summary["K"] == k].iloc[0]
        columns = {"UAV_A": "A型运输无人机数", "UAV_B": "B型运输无人机数",
                   "UAV_C": "C型运输无人机数", "BAT_A": "A型电池组数",
                   "BAT_B": "B型电池组数", "BAT_C": "C型电池组数",
                   "RELAY": "中继无人机数", "RELAY_ENERGY": "中继能源组件数"}
        totals_ok = all(int(groups[column].sum()) == int(summary[key + "_总需求"])
                        and max(0, int(groups[column].sum()) - inventory[key]) == int(summary[key + "_缺口"])
                        for key, column in columns.items())
        check(f"q4_k{k}_resource_totals", totals_ok, int(summary["总资源缺口"]))

    book = load_workbook(RUN / "results" / "最终结果提交.xlsx", read_only=True, data_only=True)
    expected_sheets = {"Q1_单点组批", "Q2_运输架次", "Q2_逐箱交付", "Q3_中继架次", "Q3_通信保障", "Q4_分区配置"}
    counts = {sheet.title: sum(any(c.value is not None for c in row) for row in sheet.iter_rows())
              for sheet in book}
    book.close()
    check("official_submission_sheets", set(counts) == expected_sheets and
          counts.get("Q1_单点组批") == 19 and counts.get("Q2_运输架次") == 22 and
          counts.get("Q2_逐箱交付") == 81 and counts.get("Q3_中继架次") == 11 and
          counts.get("Q3_通信保障") == 252 and
          counts.get("Q4_分区配置") == 6, counts)
    for name in ("data_validation", "q2_feasibility_check", "q3_feasibility_check"):
        frame = csv(name)
        column = "passed" if name == "data_validation" else "是否通过"
        check(name, bool(frame[column].all()), len(frame))

    result = {"passed": all(item["passed"] for item in CHECKS), "checks": CHECKS}
    target = HERE / "reports" / "INDEPENDENT_CHECK.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{sum(item['passed'] for item in CHECKS)}/{len(CHECKS)} independent checks passed; {target}")
    for item in CHECKS:
        if not item["passed"]:
            print("FAIL", item["check"], item["detail"])
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
