from __future__ import division

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .config import BASE_DATA_ROOT, DATA_ROOT, OFFICIAL_ROOT, TEMPLATE_XLSX


def _sheet_rows(path, sheet_name):
    wb = load_workbook(str(path), data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    wb.close()
    return rows


def _table_from_header(path, sheet_name, header_row, start_row=None):
    rows = _sheet_rows(path, sheet_name)
    header = list(rows[header_row - 1])
    while header and header[-1] is None:
        header.pop()
    data = []
    first = start_row if start_row is not None else header_row + 1
    for row in rows[first - 1:]:
        vals = list(row[:len(header)])
        if all(v is None for v in vals):
            continue
        data.append(vals)
    return pd.DataFrame(data, columns=header)


def load_nodes():
    path = BASE_DATA_ROOT / "调度中心与服务区.xlsx"
    center = _table_from_header(path, "数据", 2, 3).iloc[:1].copy()
    center.columns = ["node_id", "name", "lon", "lat", "ground_alt_m"]
    center["population"] = 0
    services = _table_from_header(path, "数据", 6, 7).copy()
    services.columns = ["node_id", "name", "lon", "lat", "ground_alt_m", "population"]
    nodes = pd.concat([center, services], ignore_index=True, sort=False)
    nodes["node_type"] = np.where(nodes["node_id"] == "O01", "dispatch", "service")
    nodes["work_alt_m"] = nodes["ground_alt_m"] + np.where(nodes["node_id"] == "O01", 0.0, 30.0)
    return nodes


def load_boxes():
    path = BASE_DATA_ROOT / "物资需求与配送时限.xlsx"
    boxes = _table_from_header(path, "逐箱货箱清单", 1, 2)
    boxes.columns = ["box_id", "service_id", "material", "mass_kg", "volume_m3",
                     "first_batch", "first_deadline_s", "expected_s", "priority"]
    boxes["is_first"] = boxes["first_batch"].astype(str).eq("是")
    boxes["is_medical"] = boxes["material"].astype(str).eq("医疗物资")
    boxes["hard_deadline_s"] = np.where(
        boxes["is_medical"], boxes["expected_s"],
        np.where(boxes["is_first"], boxes["first_deadline_s"], np.nan))
    for col in ("mass_kg", "volume_m3", "first_deadline_s", "expected_s", "priority", "hard_deadline_s"):
        boxes[col] = pd.to_numeric(boxes[col], errors="coerce")
    return boxes


def load_transport_data():
    path = BASE_DATA_ROOT / "运输无人机数据.xlsx"
    models = _table_from_header(path, "数据", 2, 3).iloc[:3].copy()
    models.columns = ["model", "name", "empty_mass_kg", "max_payload_kg", "max_volume_m3",
                      "cruise_speed_mps", "range_empty_m", "range_full_m", "usable_energy_kwh",
                      "reserve_fraction", "prep_fixed_s", "load_per_box_s", "handoff_base_s",
                      "handoff_per_box_s", "climb_speed_mps", "descent_speed_mps",
                      "climb_efficiency", "descent_efficiency"]
    numeric_cols = [c for c in models.columns if c not in ("model", "name")]
    models[numeric_cols] = models[numeric_cols].apply(pd.to_numeric)
    models["reserve_fraction"] = models["reserve_fraction"] / 100.0

    rows = _sheet_rows(path, "数据")
    drones = pd.DataFrame([list(r[:3]) for r in rows[8:16]],
                          columns=["drone_id", "model", "initial_node"])
    batteries = pd.DataFrame([list(r[:3]) for r in rows[19:22]],
                             columns=["model", "count", "full_charge_s"])
    batteries[["count", "full_charge_s"]] = batteries[["count", "full_charge_s"]].apply(pd.to_numeric)
    return models, drones, batteries


def load_relay_data():
    path = BASE_DATA_ROOT / "中继无人机数据.xlsx"
    model = _table_from_header(path, "数据", 2, 3).iloc[:1].copy()
    model.columns = ["model", "name", "empty_mass_kg", "module_mass_kg", "takeoff_mass_kg",
                     "cruise_speed_mps", "cruise_power_kw", "usable_energy_kwh", "reserve_fraction",
                     "prep_fixed_s", "link_time_s", "turnaround_s", "climb_speed_mps",
                     "descent_speed_mps", "climb_efficiency", "descent_efficiency",
                     "hover_power_kw", "comm_power_kw", "max_hover_agl_m"]
    numeric_cols = [c for c in model.columns if c not in ("model", "name")]
    model[numeric_cols] = model[numeric_cols].apply(pd.to_numeric)
    model["reserve_fraction"] = model["reserve_fraction"] / 100.0
    rows = _sheet_rows(path, "数据")
    relays = pd.DataFrame([list(r[:3]) for r in rows[6:8]],
                          columns=["relay_id", "model", "initial_node"])
    components = pd.DataFrame([list(rows[11][:3])], columns=["model", "count", "full_charge_s"])
    components[["count", "full_charge_s"]] = components[["count", "full_charge_s"]].apply(pd.to_numeric)
    return model, relays, components


def load_link_parameters():
    path = BASE_DATA_ROOT / "通信链路参数.xlsx"
    rows = _sheet_rows(path, "数据")
    records = []
    for row in rows[2:16]:
        records.append({"category": row[0], "name": row[1], "symbol": row[3], "value": float(row[4])})
    return pd.DataFrame(records)


def load_all():
    models, drones, batteries = load_transport_data()
    relay_model, relays, components = load_relay_data()
    return {
        "nodes": load_nodes(),
        "boxes": load_boxes(),
        "transport_models": models,
        "transport_drones": drones,
        "transport_batteries": batteries,
        "relay_model": relay_model,
        "relay_drones": relays,
        "relay_components": components,
        "link_parameters": load_link_parameters(),
    }


def build_inventory():
    records = []
    source_files = sorted(path for path in DATA_ROOT.rglob("*") if path.is_file())
    if TEMPLATE_XLSX.is_file():
        source_files.append(TEMPLATE_XLSX)
    for path in source_files:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            while True:
                chunk = stream.read(1024 * 1024)
                if not chunk:
                    break
                digest.update(chunk)
        records.append({
            "relative_path": (str(path.relative_to(OFFICIAL_ROOT))
                              if OFFICIAL_ROOT in path.parents else path.name),
            "suffix": path.suffix.lower(),
            "size_bytes": path.stat().st_size,
            "sha256": digest.hexdigest(),
            "role": "official_source",
            "read_only": True,
        })
    return pd.DataFrame(records)


def validate_inventory(data):
    checks = []
    checks.append(("dispatch_center_O01", int((data["nodes"]["node_id"] == "O01").sum()) == 1))
    checks.append(("service_zone_count_15", int((data["nodes"]["node_type"] == "service").sum()) == 15))
    checks.append(("box_count_80", len(data["boxes"]) == 80))
    checks.append(("box_id_unique", data["boxes"]["box_id"].is_unique))
    checks.append(("transport_models_ABC", set(data["transport_models"]["model"]) == set(["A", "B", "C"])))
    checks.append(("transport_drone_count_8", len(data["transport_drones"]) == 8))
    checks.append(("relay_drone_count_2", len(data["relay_drones"]) == 2))
    checks.append(("dem_mat_exists", (DATA_ROOT / "镇龙乡地理空间数据" / "镇龙乡及周边地理数据" /
                                      "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.mat").exists()))
    failed = [name for name, ok in checks if not ok]
    if failed:
        raise ValueError("数据完整性检查失败: " + ", ".join(failed))
    return pd.DataFrame([{"check": name, "passed": bool(ok)} for name, ok in checks])
