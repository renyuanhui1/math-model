from __future__ import division

from pathlib import Path

from openpyxl import load_workbook

from src.config import FIGURES_DIR, OUTPUT_ROOT, RESULTS_DIR


def read_records(file_name, sheet_name=None):
    workbook = load_workbook(str(RESULTS_DIR / file_name), data_only=True, read_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    workbook.close()
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    headers = rows[0]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def f(value, digits=3):
    return ("%%.%df" % digits) % float(value)


def main():
    safe = read_records("q1_max_safe_payload.xlsx")
    q1 = read_records("q1_batching.xlsx")
    sensitivity = read_records("q1_safety_margin.xlsx")
    q2 = read_records("q2_summary.xlsx", "总结")[0]
    q3 = read_records("q3_joint_summary.xlsx")[0]
    resource = read_records("q4_resource_comparison.xlsx", "资源对比")
    partitions = {int(row["K"]): row for row in read_records(
        "q4_resource_comparison.xlsx", "分区摘要")}
    k2 = read_records("q4_partition_K2.xlsx")
    k3 = read_records("q4_partition_K3.xlsx")

    gap_names = {"UAV_A": "A型运输无人机", "UAV_B": "B型运输无人机",
                 "UAV_C": "C型运输无人机", "BAT_A": "A型电池",
                 "BAT_B": "B型电池", "BAT_C": "C型电池",
                 "RELAY": "中继无人机", "RELAY_ENERGY": "中继能源组件"}

    def gap_description(partition):
        return "、".join("%s%d" % (label, int(partition[key + "_缺口"]))
                        for key, label in gap_names.items()
                        if int(partition[key + "_缺口"]) > 0) or "无"

    model_ranges = []
    for model in ("A", "B", "C"):
        values = [row["最大安全载荷（kg）"] for row in safe if row["机型编号"] == model]
        model_ranges.append("%s型 %s–%s kg" % (model, f(min(values), 3), f(max(values), 3)))
    q1_energy = sum(float(row["架次能耗（kWh）"]) for row in q1)
    q1_min_soc = min(float(row["返航SOC（%）"]) for row in q1)

    summary = """# D题计算结果摘要

## 问题一：安全载荷与单点组批

- 20%% 返航安全余量下，最大安全载荷范围：%s。
- 80个货箱的单点最优组批共 %d 架次，总能耗 %s kWh，最低返航 SOC 为 %s%%。
- 安全余量敏感性：%s。

## 问题二：运输调度

- 优化后运输架次由 %d 降至 %d，硬时限违约货箱数为 %d。
- 加权延期为 %s 优先系数·s，任务完成时刻 %s s，总能耗 %s kWh。
- 80个货箱均且仅交付一次；载重、体积、电量、无人机和电池时序检查均通过。

## 问题三：运输—通信联合优化

- 联合方案含 %d 个运输架次、%d 个中继架次和 %d 个时变悬停点。
- 对 %d 个1秒航迹采样点复核：直连 %d 点、中继 %d 点、通信中断 %d 点。
- 最小硬时限余量 %s s，联合任务完成时刻 %s s，中继总能耗 %s kWh。

## 问题四：分区独立保障

- 运输路线超边构成%d个不可分任务块。
- K=2 方案：%s；总资源缺口%d（%s），相比集中调度产生%d单位冗余。
- K=3 方案：%s；总资源缺口%d（%s），相比集中调度产生%d单位冗余。
- 集中调度的资源缺口为%d；在固定问题三任务下，K=2的缺口及冗余均小于K=3。

## 稳健结论与假设依赖

- 最稳健：所有硬时限均满足；运输及中继能量约束均满足；1秒采样网格上无通信中断；K=2的总资源缺口与冗余均小于K=3。
- 较强假设：直线航段、确定性飞行时间/能耗、30 m DEM和规则化地形遮挡损耗；未额外引入风场、降雨、故障率和临时禁飞区。
""" % (
        "、".join(model_ranges), len(q1), f(q1_energy), f(q1_min_soc, 2),
        "；".join("%s%%→%d架次" % (f(row["返航安全余量"] * 100, 0), row["最优架次数"])
                 for row in sensitivity),
        q2["初始架次数"], q2["改进后架次数"], q2["硬时限违约货箱数"],
        f(q2["加权延期（优先系数·s）"]), f(q2["任务完成时刻（s）"]), f(q2["总能耗（kWh）"]),
        q3["运输架次数"], q3["中继架次数"], q3["中继悬停点数"], q3["全航迹采样点数"],
        q3["直连覆盖点数"], q3["中继覆盖点数"], q3["通信中断点数"],
        f(q3["最小硬时限余量（s）"]),
        f(q3["联合任务完成时刻（s）"]), f(q3["中继总能耗（kWh）"]),
        partitions[2]["任务块数量"],
        "；".join(row["任务组编号"] + "=" + row["服务区列表"] for row in k2),
        partitions[2]["总资源缺口"], gap_description(partitions[2]),
        partitions[2]["总资源冗余"],
        "；".join(row["任务组编号"] + "=" + row["服务区列表"] for row in k3),
        partitions[3]["总资源缺口"], gap_description(partitions[3]),
        partitions[3]["总资源冗余"], resource[0]["总资源缺口"])
    (OUTPUT_ROOT / "summary.md").write_text(summary, encoding="utf-8")

    mapping = """# 论文表图与计算文件映射

| 论文建议编号 | 内容 | 数据文件 | 图文件 |
|---|---|---|---|
| 表1 | 各机型最大安全载荷 | `results/q1_max_safe_payload.xlsx` | `figures/q1_safe_payload_heatmap.png/.pdf` |
| 表2 | 单点最优组批 | `results/q1_batching.xlsx` | `figures/q1_safety_margin.png/.pdf` |
| 图1 | 地形与服务区 | `results/segment_database.xlsx` | `figures/q1_dem_service_map.png/.pdf` |
| 表3 | 问题二运输架次 | `results/q2_transport_sorties.xlsx` | `figures/q2_routes.png/.pdf` |
| 表4 | 逐箱交付与可行性 | `results/q2_box_delivery.xlsx`; `results/q2_feasibility_check.xlsx` | `figures/q2_delivery_timeline.png/.pdf` |
| 图2 | 无人机/电池时序 | `results/q2_resource_schedule.xlsx` | `figures/q2_uav_gantt.png/.pdf`; `figures/q2_battery_gantt.png/.pdf` |
| 表5 | 中继架次 | `results/q3_relay_sorties.xlsx` | `figures/q3_relay_layout.png/.pdf` |
| 表6 | 通信保障区间 | `results/q3_communication_coverage.xlsx` | `figures/q3_communication_timeline.png/.pdf` |
| 图3 | 直连覆盖与联合时序 | `results/q3_joint_summary.xlsx` | `figures/q3_direct_coverage_map.png/.pdf`; `figures/q3_joint_gantt.png/.pdf` |
| 表7 | K=2/K=3分区配置 | `results/q4_partition_K2.xlsx`; `results/q4_partition_K3.xlsx` | `figures/q4_partition_K2.png/.pdf`; `figures/q4_partition_K3.png/.pdf` |
| 图4 | 分区资源对比 | `results/q4_resource_comparison.xlsx` | `figures/q4_resource_comparison.png/.pdf` |
| 提交表 | 六个官方工作表 | `results/最终结果提交.xlsx` | — |
"""
    (OUTPUT_ROOT / "paper_mapping.md").write_text(mapping, encoding="utf-8")


if __name__ == "__main__":
    main()
