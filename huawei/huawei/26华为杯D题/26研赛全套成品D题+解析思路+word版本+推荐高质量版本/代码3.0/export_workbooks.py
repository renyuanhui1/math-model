from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import RESULTS_DIR, TABLE_DIR, TEMPLATE_XLSX


SPECS = [
    ("data_inventory.xlsx", [("data_inventory", "数据清单"), ("data_validation", "完整性检查")]),
    ("segment_database.xlsx", [("segment_database", "航段数据库")]),
    ("q1_max_safe_payload.xlsx", [("q1_max_safe_payload", "最大安全载荷")]),
    ("q1_batching.xlsx", [("q1_batching", "单点组批")]),
    ("q1_safety_margin.xlsx", [("q1_safety_margin", "安全余量敏感性")]),
    ("q2_transport_sorties.xlsx", [("q2_transport_sorties", "运输架次")]),
    ("q2_box_delivery.xlsx", [("q2_box_delivery", "逐箱交付")]),
    ("q2_resource_schedule.xlsx", [("q2_resource_schedule", "资源时序")]),
    ("q2_summary.xlsx", [("q2_summary", "总结"), ("q2_optimization_history", "优化过程")]),
    ("q2_feasibility_check.xlsx", [("q2_feasibility_check", "可行性检查")]),
    ("q3_direct_link_intervals.xlsx", [("q3_direct_link_intervals", "直连区间")]),
    ("q3_transport_schedule.xlsx", [("q3_transport_schedule", "运输架次"),
                                     ("q3_box_delivery", "逐箱交付"),
                                     ("q3_resource_schedule", "联合资源时序")]),
    ("q3_relay_sorties.xlsx", [("q3_relay_sorties", "中继架次")]),
    ("q3_communication_coverage.xlsx", [("q3_communication_coverage", "通信保障")]),
    ("q3_joint_summary.xlsx", [("q3_joint_summary", "联合优化总结"),
                                ("q3_optimization_history", "联合搜索过程"),
                                ("q3_feasibility_check", "可行性检查")]),
    ("q4_partition_K2.xlsx", [("q4_partition_K2", "K2分区")]),
    ("q4_partition_K3.xlsx", [("q4_partition_K3", "K3分区")]),
    ("q4_resource_comparison.xlsx", [("q4_resource_comparison", "资源对比"),
                                      ("q4_summary", "分区摘要")]),
]

FINAL_MAPPINGS = [
    ("Q1_单点组批", "q1_batching", ["架次编号", "服务区编号", "机型编号", "货箱编号列表",
                                      "总质量（kg）", "总体积（m³）", "往返时间（s）", "架次能耗（kWh）", "返航SOC（%）"]),
    ("Q2_运输架次", "q2_transport_sorties", ["架次编号", "无人机编号", "机型编号", "电池编号",
                                                   "开始时刻（s）", "访问服务区顺序", "返回O01时刻（s）", "架次能耗（kWh）"]),
    ("Q2_逐箱交付", "q2_box_delivery", ["货箱编号", "架次编号", "服务区编号", "交付完成时刻（s）"]),
    ("Q3_中继架次", "q3_relay_sorties", ["中继架次编号", "中继无人机编号", "能源组件编号", "开始时刻（s）",
                                              "悬停经度（°）", "悬停纬度（°）", "悬停海拔（m）", "建链完成时刻（s）",
                                              "服务结束时刻（s）", "返回O01时刻（s）", "架次能耗（kWh）"]),
    ("Q3_通信保障", "q3_communication_coverage", ["运输架次编号", "通信阶段", "开始时刻（s）",
                                                           "结束时刻（s）", "保障方式", "中继架次编号"]),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="1F2937")
THIN = Side(style="thin", color="D9E2F3")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def read_table(name):
    return pd.read_csv(TABLE_DIR / (name + ".csv"), encoding="utf-8-sig")


def clean_value(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if hasattr(value, "item"):
        return value.item()
    return value


def clear_sheet(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


def write_frame(ws, frame):
    clear_sheet(ws)
    matrix = [list(frame.columns)] + frame.values.tolist()
    for r_idx, row in enumerate(matrix, 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, clean_value(value))
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(r_idx == 1))
            if r_idx == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = BODY_FONT
                title = str(matrix[0][c_idx - 1])
                if re.search(r"s）|kWh|经度|纬度|高程|能耗|质量|体积|距离|时刻|时间", title):
                    cell.number_format = "0.000"
                elif re.search(r"SOC|%|余量|系数", title):
                    cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for col_idx, title in enumerate(matrix[0], 1):
        values = [str(title)] + [str(clean_value(v) or "") for v in frame.iloc[:249, col_idx - 1].tolist()]
        width = min(42, max(10, max(len(value) for value in values) * 1.08))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def save_book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for table_name, sheet_name in sheets:
        write_frame(wb.create_sheet(sheet_name), read_table(table_name))
    wb.save(path)


def export_final_submission():
    wb = load_workbook(TEMPLATE_XLSX)
    for sheet_name, table_name, columns in FINAL_MAPPINGS:
        frame = read_table(table_name).loc[:, columns]
        write_frame(wb[sheet_name], frame)
    q4_columns = ["K（2或3）", "任务组编号", "服务区列表", "A型运输无人机数", "B型运输无人机数",
                  "C型运输无人机数", "A型电池组数", "B型电池组数", "C型电池组数", "中继无人机数", "中继能源组件数"]
    q4 = pd.concat([read_table("q4_partition_K2"), read_table("q4_partition_K3")], ignore_index=True)
    write_frame(wb["Q4_分区配置"], q4.loc[:, q4_columns])
    wb.save(RESULTS_DIR / "最终结果提交.xlsx")


def main():
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    for file_name, sheets in SPECS:
        save_book(RESULTS_DIR / file_name, sheets)
    export_final_submission()
    print("exported %d workbooks with openpyxl fallback" % (len(SPECS) + 1))


if __name__ == "__main__":
    main()
